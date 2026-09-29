"""多格式文档加载器：PDF/Word/Markdown/网页/纯文本。"""
import re
from pathlib import Path

import pdfplumber
import requests
from bs4 import BeautifulSoup
from docx import Document


class DocumentLoader:
    """多格式文档加载器：PDF/Word/Markdown/网页/纯文本。"""

    # OCR 引擎懒加载单例（首次遇到无文字层的扫描件时才初始化）
    _ocr_engine = None
    OCR_RENDER_DPI = 200  # 扫描件页面转图片的渲染分辨率

    @classmethod
    def _get_ocr_engine(cls):
        """懒加载 RapidOCR 引擎，未安装时返回 None。"""
        if cls._ocr_engine is None:
            try:
                from rapidocr_onnxruntime import RapidOCR
                cls._ocr_engine = RapidOCR()
            except ImportError:
                cls._ocr_engine = False  # 标记不可用，避免反复尝试导入
        return cls._ocr_engine or None

    @classmethod
    def _ocr_page(cls, page) -> str:
        """对无文字层的 PDF 页面做 OCR，返回识别出的文本。"""
        engine = cls._get_ocr_engine()
        if engine is None:
            print(f"⚠️  检测到扫描页但未安装 OCR，请执行: pip install rapidocr-onnxruntime")
            return ""
        img = page.to_image(resolution=cls.OCR_RENDER_DPI)
        pil_img = img.original  # PIL Image，RapidOCR 支持 ndarray 输入
        import numpy as np
        result, _ = engine(np.array(pil_img))
        if not result:
            return ""
        # result 格式: [[bbox, text, confidence], ...]，按识别顺序拼接
        return "\n".join(line[1] for line in result)

    @staticmethod
    def load_pdf(path: str) -> tuple[str, dict]:
        """加载 PDF，提取文本和表格（转 Markdown）；无文字层的扫描页自动走 OCR 兜底。"""
        content_parts = []
        page_map = []  # 记录每个内容块的页码
        ocr_pages = []  # 走了 OCR 的页码列表

        with pdfplumber.open(path) as pdf:
            total_pages = len(pdf.pages)
            for page_num, page in enumerate(pdf.pages, start=1):
                page_content = []

                # 提取表格并转 Markdown
                tables = page.extract_tables()
                for table in tables:
                    md_table = DocumentLoader._table_to_markdown(table)
                    if md_table:
                        page_content.append(md_table)

                # 提取普通文本
                text = page.extract_text()
                if text:
                    page_content.append(text)

                # 文字层为空但有图片 → 扫描件，走 OCR 兜底
                if not text and page.images:
                    ocr_text = DocumentLoader._ocr_page(page)
                    if ocr_text.strip():
                        page_content.append(ocr_text.strip())
                        ocr_pages.append(page_num)

                if page_content:
                    combined = "\n\n".join(page_content)
                    content_parts.append(combined)
                    page_map.append({"page": page_num, "length": len(combined)})

        full_content = "\n\n".join(content_parts)
        metadata = {"page_map": page_map, "total_pages": total_pages}
        if ocr_pages:
            metadata["ocr_pages"] = ocr_pages
        return full_content, metadata

    @staticmethod
    def load_docx(path: str) -> tuple[str, dict]:
        """加载 Word 文档，保留标题层级和表格。"""
        doc = Document(path)
        content_parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                # 根据标题级别转 Markdown
                if para.style.name.startswith("Heading"):
                    level = int(para.style.name.replace("Heading", "").strip() or "1")
                    content_parts.append(f"{'#' * level} {para.text}\n\n")
                else:
                    content_parts.append(f"{para.text}\n\n")

        # 提取表格
        for table in doc.tables:
            md_table = DocumentLoader._docx_table_to_markdown(table)
            if md_table:
                content_parts.append(md_table)
                content_parts.append("\n\n")

        return "".join(content_parts).strip(), {}

    @staticmethod
    def load_markdown(path: str) -> tuple[str, dict]:
        """加载 Markdown 文件。"""
        with open(path, encoding="utf-8") as f:
            return f.read(), {}

    @staticmethod
    def load_text(path: str) -> tuple[str, dict]:
        """加载纯文本文件。"""
        with open(path, encoding="utf-8") as f:
            return f.read(), {}

    @staticmethod
    def load_html(url_or_path: str) -> tuple[str, dict]:
        """加载网页或 HTML 文件，提取正文内容。"""
        if url_or_path.startswith(("http://", "https://")):
            resp = requests.get(url_or_path, timeout=30)
            resp.raise_for_status()
            html_content = resp.text
        else:
            with open(url_or_path, encoding="utf-8") as f:
                html_content = f.read()

        soup = BeautifulSoup(html_content, "html.parser")

        # 移除脚本和样式
        for script in soup(["script", "style", "nav", "footer", "header"]):
            script.decompose()

        # 提取标题
        title = soup.title.string.strip() if soup.title else ""

        # 提取正文（优先 article/main 标签）
        main_content = soup.find("article") or soup.find("main") or soup.find("body")

        if main_content:
            # 保留标题层级
            for h in main_content.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
                level = int(h.name[1])
                h.replace_with(f"\n{'#' * level} {h.get_text(strip=True)}\n\n")

            # 保留表格
            for table in main_content.find_all("table"):
                md_table = DocumentLoader._html_table_to_markdown(table)
                table.replace_with(f"\n{md_table}\n\n")

            # 保留代码块
            for code in main_content.find_all(["pre", "code"]):
                code_text = code.get_text()
                code.replace_with(f"\n```\n{code_text}\n```\n\n")

            text = main_content.get_text(separator="\n")
        else:
            text = soup.get_text(separator="\n")

        metadata = {"title": title, "source_url": url_or_path}
        return text, metadata

    @staticmethod
    def _table_to_markdown(table: list[list]) -> str:
        """将 pdfplumber 提取的表格转为 Markdown 格式。"""
        if not table or not table[0]:
            return ""

        # 清理单元格内容
        cleaned = []
        for row in table:
            cleaned_row = [str(cell).strip().replace("\n", " ") if cell else "" for cell in row]
            cleaned.append(cleaned_row)

        # 生成 Markdown 表格
        lines = []
        lines.append("| " + " | ".join(cleaned[0]) + " |")
        lines.append("| " + " | ".join(["---"] * len(cleaned[0])) + " |")
        for row in cleaned[1:]:
            while len(row) < len(cleaned[0]):
                row.append("")
            lines.append("| " + " | ".join(row[:len(cleaned[0])]) + " |")

        return "\n".join(lines)

    @staticmethod
    def _docx_table_to_markdown(table) -> str:
        """将 python-docx 的表格转为 Markdown 格式。"""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(cells)

        if not rows:
            return ""

        lines = []
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
        for row in rows[1:]:
            lines.append("| " + " | ".join(row) + " |")

        return "\n".join(lines)

    @staticmethod
    def _html_table_to_markdown(table) -> str:
        """将 HTML 表格转为 Markdown 格式。"""
        rows = []
        for tr in table.find_all("tr"):
            cells = [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]
            if cells:
                rows.append(cells)

        if not rows:
            return ""

        lines = []
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
        for row in rows[1:]:
            lines.append("| " + " | ".join(row) + " |")

        return "\n".join(lines)

    @classmethod
    def load(cls, path: str) -> tuple[str, dict]:
        """根据文件扩展名自动选择加载器。"""
        ext = Path(path).suffix.lower()
        if ext == ".pdf":
            return cls.load_pdf(path)
        elif ext == ".docx":
            return cls.load_docx(path)
        elif ext == ".md":
            return cls.load_markdown(path)
        elif ext in (".html", ".htm"):
            return cls.load_html(path)
        elif path.startswith(("http://", "https://")):
            return cls.load_html(path)
        else:
            return cls.load_text(path)
